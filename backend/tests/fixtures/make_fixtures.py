from __future__ import annotations

import io
import zipfile
from pathlib import Path

import docx
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas


def get_tiny_pdf(
    text: str = "VaultRAG tiny test PDF document content for testing ingestion pipeline.",
) -> bytes:
    """Generate a tiny valid PDF with extractable text via reportlab."""
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=letter)
    c.drawString(100, 750, text)
    c.showPage()
    c.save()
    return buf.getvalue()


def get_tiny_docx(
    text: str = "VaultRAG tiny test DOCX document paragraph for testing ingestion pipeline.",
) -> bytes:
    """Generate a tiny valid DOCX with paragraphs and a table via python-docx."""
    buf = io.BytesIO()
    doc = docx.Document()
    doc.add_paragraph(text)
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Header Col A"
    table.rows[0].cells[1].text = "Header Col B"
    doc.save(buf)
    return buf.getvalue()


def get_tiny_txt(
    text: str = "VaultRAG tiny test TXT plain text document for testing ingestion pipeline.",
) -> bytes:
    """Generate UTF-8 encoded plain text bytes."""
    return text.encode("utf-8")


DEFAULT_MD_TEXT = (
    "# VaultRAG Markdown\n\nThis is a tiny test MD document for testing ingestion pipeline."
)


def get_tiny_md(
    text: str = DEFAULT_MD_TEXT,
) -> bytes:
    """Generate UTF-8 encoded markdown bytes."""
    return text.encode("utf-8")


def get_corrupt_pdf() -> bytes:
    """Generate a corrupt PDF: starts with %PDF header followed by corrupt binary garbage."""
    return b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\nGARBAGE_NOT_A_VALID_PDF_TRAILER_OR_CATALOG_AT_ALL"


def get_fake_pdf() -> bytes:
    """Generate a fake PDF: plain text without %PDF header."""
    return (
        b"This is plain text pretending to be a PDF file. It does not start with %PDF magic bytes."
    )


def get_zipbomb_docx() -> bytes:
    """Generate a simulated DOCX zip bomb (>50MB uncompressed) to test zip-bomb guard."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        zf.writestr(
            "word/document.xml",
            b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            b'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            b"<w:body><w:p><w:r><w:t>Bomb</w:t></w:r></w:p></w:body></w:document>",
        )
        # 52 MB of zeros compressed with deflate is ~53 KB
        zf.writestr("bomb_payload.bin", b"\x00" * (52 * 1024 * 1024))
    return buf.getvalue()


def make_all_fixtures(output_dir: Path | str | None = None) -> dict[str, bytes]:
    """Generate all test fixtures and optionally write them to disk."""
    fixtures = {
        "tiny.pdf": get_tiny_pdf(),
        "tiny.docx": get_tiny_docx(),
        "tiny.txt": get_tiny_txt(),
        "tiny.md": get_tiny_md(),
        "corrupt.pdf": get_corrupt_pdf(),
        "fake.pdf": get_fake_pdf(),
        "zipbomb.docx": get_zipbomb_docx(),
    }

    if output_dir:
        out = Path(output_dir)
        out.mkdir(parents=True, exist_ok=True)
        for name, data in fixtures.items():
            (out / name).write_bytes(data)

    return fixtures


if __name__ == "__main__":
    target = Path(__file__).parent / "data"
    make_all_fixtures(target)
    print(f"Generated test fixtures in {target}")
