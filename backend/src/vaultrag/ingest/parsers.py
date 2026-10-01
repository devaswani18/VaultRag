from __future__ import annotations

import io
import zipfile

import docx
from pypdf import PdfReader

from vaultrag.errors import ValidationFailed


def _parse_pdf(data: bytes) -> list[tuple[int, str]]:
    """Validate PDF magic bytes, guard against corrupt/encrypted files and parse pages."""
    if not data.startswith(b"%PDF"):
        raise ValidationFailed("Invalid PDF file: missing %PDF header")

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                if reader.decrypt("") == 0:
                    raise ValidationFailed("Encrypted PDF files are not supported")
            except Exception:
                raise ValidationFailed("Encrypted PDF files are not supported") from None

        if len(reader.pages) > 200:
            raise ValidationFailed("PDF exceeds maximum allowed page count (200)")

        pages: list[tuple[int, str]] = []
        for idx, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            if text.strip():
                pages.append((idx, text.strip()))
        return pages
    except ValidationFailed:
        raise
    except Exception as e:
        raise ValidationFailed(f"Corrupt or invalid PDF file: {type(e).__name__}") from e


def _parse_docx(data: bytes) -> list[tuple[int, str]]:
    """Validate DOCX zip archive, guard against zip bombs (<50MB uncompressed) and parse content."""
    if not zipfile.is_zipfile(io.BytesIO(data)):
        raise ValidationFailed("Invalid DOCX file: not a valid zip archive")

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            namelist = zf.namelist()
            if "word/document.xml" not in namelist:
                raise ValidationFailed("Invalid DOCX file: missing word/document.xml")

            total_uncompressed = sum(info.file_size for info in zf.infolist())
            if total_uncompressed > 50 * 1024 * 1024:
                raise ValidationFailed(
                    "DOCX uncompressed size exceeds 50 MB limit (zip-bomb guard)"
                )
    except ValidationFailed:
        raise
    except Exception as e:
        raise ValidationFailed(f"Corrupt or invalid DOCX archive: {type(e).__name__}") from e

    try:
        doc = docx.Document(io.BytesIO(data))
        lines: list[str] = []
        for p in doc.paragraphs:
            if p.text.strip():
                lines.append(p.text.strip())
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    lines.append(" | ".join(cells))

        full_text = "\n\n".join(lines).strip()
        if not full_text:
            return []
        return [(1, full_text)]
    except Exception as e:
        raise ValidationFailed(f"Failed to parse DOCX document: {type(e).__name__}") from e


def _parse_text(data: bytes) -> list[tuple[int, str]]:
    """Decode plain text or markdown as UTF-8, allowing replace only if >99% valid."""
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError:
        decoded = data.decode("utf-8", errors="replace")
        replacement_count = decoded.count("\ufffd")
        valid_ratio = 1.0 - (replacement_count / max(1, len(decoded)))
        if valid_ratio < 0.99:
            raise ValidationFailed("File is not valid UTF-8 text") from None
        text = decoded

    text = text.strip()
    if not text:
        return []
    return [(1, text)]


def parse_document(data: bytes, ext: str) -> list[tuple[int, str]]:
    """Parse document bytes into a list of (page_number, text) tuples based on extension."""
    clean_ext = ext.lower().lstrip(".")
    if clean_ext == "pdf":
        return _parse_pdf(data)
    elif clean_ext == "docx":
        return _parse_docx(data)
    elif clean_ext in ("txt", "md"):
        return _parse_text(data)
    else:
        raise ValidationFailed(f"Unsupported document extension: {clean_ext}")
