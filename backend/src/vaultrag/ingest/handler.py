from __future__ import annotations

import io
import logging
import urllib.parse
import zipfile
from datetime import UTC, datetime
from typing import Any

import docx
from pypdf import PdfReader
from qdrant_client import models

from vaultrag.clients.dynamo import DocumentRepo, DocumentStatus, TenantRepo
from vaultrag.clients.gemini import embed_texts
from vaultrag.clients.qdrant import upsert_chunks
from vaultrag.clients.s3 import download_bytes, parse_object_key
from vaultrag.config import get_settings
from vaultrag.errors import NotFound, ValidationFailed
from vaultrag.ingest.chunker import chunk_document
from vaultrag.ingest.pipeline import DocumentBlocked, get_default_hooks, run_hooks

logger = logging.getLogger(__name__)


def _safe_error_message(exc: Exception) -> str:
    """Return safe user-facing error message without raw document or prompt text."""
    if isinstance(exc, (ValidationFailed, DocumentBlocked)):
        return exc.message
    return f"Processing failed: {type(exc).__name__}"


def _parse_pdf(data: bytes) -> list[tuple[int, str]]:
    """Validate PDF magic bytes, guard against corrupt/encrypted files and parse pages."""
    if not data.startswith(b"%PDF"):
        raise ValidationFailed("Invalid PDF file: missing %PDF header")

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                if reader.decrypt("") == 0:
                    raise ValidationFailed("Encrypted PDF files are not supported")
            except Exception:
                raise ValidationFailed("Encrypted PDF files are not supported") from None

        if len(reader.pages) > 200:
            raise ValidationFailed("PDF exceeds maximum allowed page count (200)")

        pages: list[tuple[int, str]] = []
        for idx, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            if text.strip():
                pages.append((idx, text.strip()))
        return pages
    except ValidationFailed:
        raise
    except Exception as e:
        raise ValidationFailed(f"Corrupt or invalid PDF file: {type(e).__name__}") from e


def _parse_docx(data: bytes) -> list[tuple[int, str]]:
    """Validate DOCX zip archive, guard against zip bombs (<50MB uncompressed) and parse content."""
    if not zipfile.is_zipfile(io.BytesIO(data)):
        raise ValidationFailed("Invalid DOCX file: not a valid zip archive")

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            namelist = zf.namelist()
            if "word/document.xml" not in namelist:
                raise ValidationFailed("Invalid DOCX file: missing word/document.xml")

            total_uncompressed = sum(info.file_size for info in zf.infolist())
            if total_uncompressed > 50 * 1024 * 1024:
                raise ValidationFailed(
                    "DOCX uncompressed size exceeds 50 MB limit (zip-bomb guard)"
                )
    except ValidationFailed:
        raise
    except Exception as e:
        raise ValidationFailed(f"Corrupt or invalid DOCX archive: {type(e).__name__}") from e

    try:
        doc = docx.Document(io.BytesIO(data))
        lines: list[str] = []
        for p in doc.paragraphs:
            if p.text.strip():
                lines.append(p.text.strip())
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    lines.append(" | ".join(cells))

        full_text = "\n\n".join(lines).strip()
        if not full_text:
            return []
        return [(1, full_text)]
    except Exception as e:
        raise ValidationFailed(f"Failed to parse DOCX document: {type(e).__name__}") from e


def _parse_text(data: bytes) -> list[tuple[int, str]]:
    """Decode plain text or markdown as UTF-8, allowing replace only if >99% valid."""
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError:
        decoded = data.decode("utf-8", errors="replace")
        replacement_count = decoded.count("\ufffd")
        valid_ratio = 1.0 - (replacement_count / max(1, len(decoded)))
        if valid_ratio < 0.99:
            raise ValidationFailed("File is not valid UTF-8 text") from None
        text = decoded

    text = text.strip()
    if not text:
        return []
    return [(1, text)]


def parse_document(data: bytes, ext: str) -> list[tuple[int, str]]:
    """Parse document bytes into a list of (page_number, text) tuples based on extension."""
    clean_ext = ext.lower().lstrip(".")
    if clean_ext == "pdf":
        return _parse_pdf(data)
    elif clean_ext == "docx":
        return _parse_docx(data)
    elif clean_ext in ("txt", "md"):
        return _parse_text(data)
    else:
        raise ValidationFailed(f"Unsupported document extension: {clean_ext}")


def process_record(record: dict[str, Any]) -> None:
    """Process a single S3 ObjectCreated event record."""
    s3_info = record.get("s3", {})
    raw_key = s3_info.get("object", {}).get("key", "")
    if not raw_key:
        logger.warning("Record missing s3.object.key, skipping")
        return

    # S3 event object keys are URL-encoded
    key = urllib.parse.unquote_plus(raw_key)

    # 1. Parse object key (tenant_id, doc_id, ext). Ignore and log any non-matching keys.
    try:
        tenant_id, doc_id, ext = parse_object_key(key)
    except ValidationFailed:
        logger.warning("Ignoring S3 object with invalid key structure: %s", key)
        return

    # 2. Load document record by (tenant_id, doc_id). If missing: ignore and log.
    repo = DocumentRepo()
    try:
        doc = repo.get(tenant_id, doc_id)
    except NotFound:
        logger.warning(
            "Document record not found for tenant_id=%s, doc_id=%s. Ignoring event.",
            tenant_id,
            doc_id,
        )
        return
    except Exception as e:
        logger.error(
            "Unexpected error loading document doc_id=%s for tenant_id=%s: %s",
            doc_id,
            tenant_id,
            type(e).__name__,
        )
        return

    # Idempotency check: if status != PENDING_UPLOAD, skip duplicate event
    curr_status = doc.get("status")
    if curr_status != DocumentStatus.PENDING_UPLOAD.value:
        logger.info(
            "Document doc_id=%s status is '%s' (not PENDING_UPLOAD). Skipping duplicate event.",
            doc_id,
            curr_status,
        )
        return

    # 3-10. Processing pipeline with safe exception handling
    try:
        # 3. Set PROCESSING and verify size constraints
        repo.update_status(tenant_id, doc_id, DocumentStatus.PROCESSING)

        settings = get_settings()
        max_bytes = settings.max_upload_mb * 1024 * 1024
        declared_size = int(doc.get("size_bytes", 0))

        # 4. Download bytes (bounded) and validate magic bytes
        bucket = s3_info.get("bucket", {}).get("name")
        data = download_bytes(key, max_bytes=max_bytes, bucket=bucket)
        actual_size = len(data)

        if actual_size <= 0:
            raise ValidationFailed("Uploaded object is empty")
        if actual_size > max_bytes:
            raise ValidationFailed(
                f"Uploaded object size {actual_size} exceeds max {max_bytes} bytes"
            )

        # Verify object matches declared size within tolerance
        tolerance = max(1024, int(declared_size * 0.05))
        if abs(actual_size - declared_size) > tolerance:
            raise ValidationFailed("Uploaded size does not match declared size within tolerance")

        # 5. Parse -> list of (page, text)
        pages = parse_document(data, ext)

        total_text_len = sum(len(text.strip()) for _, text in pages)
        if total_text_len < 20:
            raise ValidationFailed("no extractable text (scanned documents are not supported yet)")

        # 6. Chunk recursively (target 1000, overlap 150, min 100, deterministic uuid5 IDs)
        chunks = chunk_document(
            pages=pages,
            tenant_id=tenant_id,
            doc_id=doc_id,
            target_size=1000,
            overlap=150,
            min_size=100,
        )
        if not chunks:
            raise ValidationFailed("no extractable text (scanned documents are not supported yet)")

        # 7. Run hook pipeline
        tenant_repo = TenantRepo()
        try:
            tenant_record = tenant_repo.get(tenant_id)
            tenant_settings = tenant_record.get("settings", {})
        except Exception:
            tenant_settings = {}

        # ----------------------------------------------------------------------
        # NOTE ON HOOK REGISTRATION:
        # In Stage 8, get_default_hooks() returns an empty list.
        # Stage 9: PII guard hook will be registered to detect/redact PII.
        # Stage 10: Prompt injection guard hook will be registered to scan chunks.
        # ----------------------------------------------------------------------
        hooks = get_default_hooks()
        contexts = run_hooks(chunks, hooks, tenant_settings)
        active_contexts = [ctx for ctx in contexts if not ctx.is_dropped and ctx.text.strip()]

        if not active_contexts:
            raise ValidationFailed("All document chunks were dropped during processing")

        # 8. Embed all remaining chunks (RETRIEVAL_DOCUMENT) in batches
        texts_to_embed = [ctx.text for ctx in active_contexts]
        embeddings = embed_texts(texts_to_embed, task_type="RETRIEVAL_DOCUMENT")

        # 9. Upsert to Qdrant with complete schema payload
        now_iso = datetime.now(UTC).isoformat()
        points: list[models.PointStruct] = []
        for ctx, emb in zip(active_contexts, embeddings, strict=True):
            payload = {
                "tenant_id": tenant_id,
                "doc_id": doc_id,
                "chunk_index": ctx.chunk.index,
                "text": ctx.text,
                "filename": doc.get("filename", ""),
                "page": ctx.chunk.page,
                "visibility": doc.get("visibility", "tenant"),
                "allowed_roles": doc.get("allowed_roles", []),
                "allowed_users": doc.get("allowed_users", []),
                "owner_user_id": doc.get("owner_user_id", ""),
                "pii_found": False,
                "pii_types": [],
                "injection_risk": "low",
                "created_at": now_iso,
            }
            if ctx.payload:
                payload.update(ctx.payload)

            points.append(
                models.PointStruct(
                    id=ctx.chunk.chunk_id,
                    vector=emb,
                    payload=payload,
                )
            )

        upsert_chunks(points)

        # 10. Update document to READY with chunk_count
        repo.update_status(
            tenant_id,
            doc_id,
            DocumentStatus.READY,
            extra_fields={"chunk_count": len(points)},
        )
        logger.info(
            "Document ingestion completed: tenant_id=%s, doc_id=%s, chunk_count=%d",
            tenant_id,
            doc_id,
            len(points),
        )

    except DocumentBlocked as e:
        # Quarantine blocked document
        logger.warning(
            "Document blocked by security policy: tenant_id=%s, doc_id=%s",
            tenant_id,
            doc_id,
        )
        safe_msg = _safe_error_message(e)
        try:
            repo.update_status(tenant_id, doc_id, DocumentStatus.QUARANTINED, error=safe_msg)
        except Exception:
            logger.exception(
                "Failed to update document to QUARANTINED for tenant_id=%s, doc_id=%s",
                tenant_id,
                doc_id,
            )
    except Exception as e:
        # Safe short error string, log exception type only (NEVER log document text)
        logger.error(
            "Document ingestion failed: tenant_id=%s, doc_id=%s, error_type=%s",
            tenant_id,
            doc_id,
            type(e).__name__,
        )
        safe_msg = _safe_error_message(e)
        try:
            repo.update_status(tenant_id, doc_id, DocumentStatus.FAILED, error=safe_msg)
        except Exception:
            logger.exception(
                "Failed to update document to FAILED for tenant_id=%s, doc_id=%s",
                tenant_id,
                doc_id,
            )


def handler(event: dict[str, Any], context: Any = None) -> dict[str, Any]:
    """AWS Lambda entry point for S3 ObjectCreated events."""
    records = event.get("Records", [])
    logger.info("Ingest handler received %d S3 record(s)", len(records))

    for record in records:
        process_record(record)

    return {"statusCode": 200, "body": "OK"}
